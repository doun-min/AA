"""엑셀 데이터 전처리: display=no 제품을 전체 시트에서 제거.

`product_spec_basics` 시트에서 `display` 값이 no 인 행의 `model_code` 를 모아,
워크북의 모든 시트에서 그 model_code 에 해당하는 행을 삭제한 새 파일을 만든다.

큰 워크북(수십만 행)을 다루므로 read-only 스트리밍으로 읽어 write-only 로
새로 쓴다. 값/헤더는 그대로 보존되지만 셀 서식·수식·열 너비 등 스타일은
보존되지 않는다(데이터 덤프 전처리 용도).

사용법:
    python drop_hidden_products.py <excel_file> [옵션]

옵션:
    -o, --output PATH   결과 저장 경로 (기본: <원본>.filtered.xlsx)
    --in-place          원본 파일에 덮어쓰기 (원본은 <원본>.bak.xlsx 로 백업)
    --dry-run           저장하지 않고 무엇이 지워질지 요약만 출력
    --basics-sheet NAME 기준 시트명 (기본: product_spec_basics)
    --display-col NAME  display 컬럼명 (기본: display)
    --model-col NAME    model_code 컬럼명 (기본: model_code)
    --extra-model-cols A,B
                        model_code 와 동일하게 취급할 추가 컬럼명
                        (기본: buying_model_code,representative_model)
"""

import argparse
import os
import shutil
import sys

import openpyxl

DEFAULT_BASICS_SHEET = "product_spec_basics"
DEFAULT_DISPLAY_COL = "display"
DEFAULT_MODEL_COL = "model_code"
DEFAULT_EXTRA_MODEL_COLS = ["buying_model_code", "representative_model"]
DROP_VALUES = {"no", "n", "false", "0"}  # display 가 이 값이면 '숨김'으로 간주


def _norm(v):
    return None if v is None else str(v).strip()


def _resolve_path(name):
    if os.path.isfile(name):
        return name
    alt = os.path.join(os.path.dirname(os.path.abspath(__file__)), name)
    if os.path.isfile(alt):
        return alt
    raise SystemExit(f"[error] 파일을 찾을 수 없습니다: {name}")


def collect_hidden_model_codes(wb, basics_sheet, display_col, model_col):
    if basics_sheet not in wb.sheetnames:
        raise SystemExit(
            f"[error] 기준 시트 '{basics_sheet}' 가 없습니다. 시트: {wb.sheetnames}"
        )
    ws = wb[basics_sheet]
    rows = ws.iter_rows(values_only=True)
    header = list(next(rows, ()) or ())
    for need in (display_col, model_col):
        if need not in header:
            raise SystemExit(
                f"[error] '{basics_sheet}' 에 '{need}' 컬럼이 없습니다. 헤더: {header}"
            )
    di, mi = header.index(display_col), header.index(model_col)

    hidden, scanned = set(), 0
    for row in rows:
        scanned += 1
        dval = _norm(row[di] if di < len(row) else None)
        if dval is not None and dval.lower() in DROP_VALUES:
            mval = _norm(row[mi] if mi < len(row) else None)
            if mval:
                hidden.add(mval)
    print(
        f"[info] {basics_sheet}: {scanned}행 스캔 -> "
        f"{display_col} ∈ {sorted(DROP_VALUES)} 인 숨김 model_code {len(hidden)}개"
    )
    return hidden


def filter_workbook(src_wb, hidden, model_col_names, dst_wb=None):
    """src_wb(read-only)를 순회하며 hidden 에 해당하는 행을 제외해 dst_wb 로 복사.

    dst_wb 가 None 이면 실제 복사는 하지 않고 통계만 계산(dry-run).
    반환: [(sheet, used_cols, before, deleted), ...]
    """
    stats = []
    for name in src_wb.sheetnames:
        src = src_wb[name]
        dst = dst_wb.create_sheet(title=name) if dst_wb is not None else None

        rows = src.iter_rows(values_only=True)
        header = list(next(rows, ()) or ())
        if dst is not None:
            dst.append(header)
        cols = [i for i, h in enumerate(header) if h in model_col_names]

        before = deleted = 0
        for row in rows:
            before += 1
            hit = any(
                ci < len(row)
                and row[ci] is not None
                and str(row[ci]).strip() in hidden
                for ci in cols
            )
            if hit:
                deleted += 1
                continue
            if dst is not None:
                dst.append(list(row))
        stats.append((name, [header[i] for i in cols], before, deleted))
    return stats


def _print_report(stats):
    print()
    print(f"{'sheet':<26} {'model_cols':<34} {'before':>8} {'deleted':>8} {'after':>8}")
    print("-" * 88)
    tot_before = tot_del = 0
    for name, used, before, deleted in stats:
        tot_before += before
        tot_del += deleted
        cols_txt = ",".join(map(str, used)) if used else "(none - skipped)"
        print(f"{name:<26} {cols_txt:<34} {before:>8} {deleted:>8} {before - deleted:>8}")
    print("-" * 88)
    print(f"{'TOTAL':<26} {'':<34} {tot_before:>8} {tot_del:>8} {tot_before - tot_del:>8}")


def main(argv=None):
    ap = argparse.ArgumentParser(
        description="display=no 제품을 엑셀 전체 시트에서 제거",
        formatter_class=argparse.RawDescriptionHelpFormatter,
    )
    ap.add_argument("excel_file", help="입력 엑셀 파일 (경로 또는 스크립트 폴더 기준 파일명)")
    ap.add_argument("-o", "--output")
    ap.add_argument("--in-place", action="store_true")
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--basics-sheet", default=DEFAULT_BASICS_SHEET)
    ap.add_argument("--display-col", default=DEFAULT_DISPLAY_COL)
    ap.add_argument("--model-col", default=DEFAULT_MODEL_COL)
    ap.add_argument("--extra-model-cols", default=",".join(DEFAULT_EXTRA_MODEL_COLS))
    args = ap.parse_args(argv)

    if args.output and args.in_place:
        raise SystemExit("[error] --output 과 --in-place 는 함께 쓸 수 없습니다.")

    path = _resolve_path(args.excel_file)
    extra = [c.strip() for c in args.extra_model_cols.split(",") if c.strip()]
    model_col_names = {args.model_col, *extra}
    print(f"[info] 입력: {path}")
    print(f"[info] model_code 로 취급할 컬럼: {sorted(model_col_names)}")

    read_wb = openpyxl.load_workbook(path, read_only=True)
    try:
        hidden = collect_hidden_model_codes(
            read_wb, args.basics_sheet, args.display_col, args.model_col
        )
        if not hidden:
            print("[info] 숨김 대상이 없습니다. 변경 없이 종료.")
            return 0

        if args.dry_run:
            stats = filter_workbook(read_wb, hidden, model_col_names, dst_wb=None)
            _print_report(stats)
            print("\n[dry-run] 파일을 저장하지 않았습니다.")
            return 0

        write_wb = openpyxl.Workbook(write_only=True)
        stats = filter_workbook(read_wb, hidden, model_col_names, dst_wb=write_wb)
        _print_report(stats)
    finally:
        read_wb.close()

    if args.in_place:
        stem, ext = os.path.splitext(path)
        bak = f"{stem}.bak{ext}"
        shutil.copy2(path, bak)
        print(f"\n[info] 원본 백업: {bak}")
        out = path
    else:
        stem, ext = os.path.splitext(path)
        out = args.output or f"{stem}.filtered{ext}"
    write_wb.save(out)
    print(f"[info] 저장 완료: {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
