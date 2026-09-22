"""Merge per-site PF v5 JSON records into US, CA, and CA_FR Excel sheets."""

import argparse
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill


SITE_SHEETS = {"us": "US", "ca": "CA", "ca_fr": "CA_FR"}
BASE_HEADERS = [
    "sku",
    "model_code",
    "display_name_pf",
    "product_color_pf",
    "sort_type",
    "sorting_no",
    "sort_presence",
    "key_model_yn",
    "source_pf_url",
    "category",
    "category_url",
]


def parse_site_input(value):
    if "=" not in value:
        raise argparse.ArgumentTypeError("use SITE=JSON_PATH")
    site, raw_path = value.split("=", 1)
    site = site.strip().lower().replace("-", "_")
    if site not in SITE_SHEETS:
        raise argparse.ArgumentTypeError("SITE must be us, ca, or ca_fr")
    path = Path(raw_path.strip())
    if not path.is_file():
        raise argparse.ArgumentTypeError(f"JSON file does not exist: {path}")
    return site, path


def excel_value(value):
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    return value


def load_records(path):
    data = json.loads(path.read_text(encoding="utf-8"))
    records = data.get("records")
    if not isinstance(records, list):
        raise ValueError(f"{path}: records must be a list")
    if not all(isinstance(record, dict) for record in records):
        raise ValueError(f"{path}: every record must be an object")
    return records, data.get("metadata") or {}


def record_headers(records):
    discovered = []
    for record in records:
        for key in record:
            if key not in BASE_HEADERS and key not in discovered:
                discovered.append(key)
    return BASE_HEADERS + discovered


def replace_site_sheet(workbook, site, records, metadata, source_path):
    title = SITE_SHEETS[site]
    if title in workbook.sheetnames:
        index = workbook.sheetnames.index(title)
        workbook.remove(workbook[title])
        sheet = workbook.create_sheet(title, index)
    else:
        sheet = workbook.create_sheet(title)

    headers = record_headers(records)
    sheet.append(headers)
    for record in records:
        sheet.append([excel_value(record.get(header)) for header in headers])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column, header in enumerate(headers, 1):
        sample_lengths = [len(str(header))]
        for row in range(2, min(sheet.max_row, 201) + 1):
            sample_lengths.append(len(str(sheet.cell(row, column).value or "")))
        sheet.column_dimensions[sheet.cell(1, column).column_letter].width = min(
            max(sample_lengths) + 2, 45
        )

    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.sheet_view.showGridLines = False


def ensure_site_sheets(workbook):
    for title in SITE_SHEETS.values():
        if title not in workbook.sheetnames:
            sheet = workbook.create_sheet(title)
            sheet.append(BASE_HEADERS)
            for cell in sheet[1]:
                cell.font = Font(color="FFFFFF", bold=True)
                cell.fill = PatternFill("solid", fgColor="7F8C8D")
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions


def update_workbook(output_path, inputs):
    if output_path.exists():
        workbook = load_workbook(output_path)
    else:
        workbook = Workbook()
        workbook.remove(workbook.active)

    ensure_site_sheets(workbook)
    counts = {}
    for site, path in inputs:
        records, metadata = load_records(path)
        replace_site_sheet(workbook, site, records, metadata, path)
        counts[site] = len(records)

    ordered = [workbook[title] for title in SITE_SHEETS.values()]
    remaining = [sheet for sheet in workbook.worksheets if sheet not in ordered]
    workbook._sheets = ordered + remaining

    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary = output_path.with_suffix(output_path.suffix + ".tmp")
    workbook.save(temporary)
    workbook.close()
    temporary.replace(output_path)
    return counts


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=Path("pf_card_details_v5.xlsx"))
    parser.add_argument(
        "--input",
        action="append",
        type=parse_site_input,
        required=True,
        metavar="SITE=JSON_PATH",
    )
    args = parser.parse_args()
    counts = update_workbook(args.output, args.input)
    print(f"[done] output={args.output.resolve()} updated={counts}")


if __name__ == "__main__":
    main()
