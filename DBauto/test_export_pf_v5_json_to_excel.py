import json
from pathlib import Path
import unittest
from uuid import uuid4

from openpyxl import load_workbook

from export_pf_v5_json_to_excel import update_workbook


class ExportPFV5JSONToExcelTests(unittest.TestCase):
    def write_json(self, path, sku):
        path.write_text(json.dumps({
            "records": [{
                "sku": sku,
                "model_code": sku,
                "sort_type": "recommended",
                "sorting_no": 1,
                "key_model_yn": "Y",
            }],
            "metadata": {"created_at": "2026-09-22 12:00:00"},
        }), encoding="utf-8")

    def test_incremental_site_update_preserves_existing_sheets(self):
        token = uuid4().hex
        us_json = Path(f"_test_pf_export_us_{token}.json")
        ca_fr_json = Path(f"_test_pf_export_cafr_{token}.json")
        output = Path(f"_test_pf_export_{token}.xlsx")
        for path in (us_json, ca_fr_json, output):
            self.addCleanup(path.unlink, missing_ok=True)
        try:
            self.write_json(us_json, "US-SKU")
            self.write_json(ca_fr_json, "CAFR-SKU")

            update_workbook(output, [("us", us_json)])
            update_workbook(output, [("ca_fr", ca_fr_json)])

            workbook = load_workbook(output, read_only=True, data_only=True)
            self.assertEqual(workbook.sheetnames[:3], ["US", "CA", "CA_FR"])
            self.assertEqual(workbook["US"]["A2"].value, "US-SKU")
            self.assertEqual(workbook["CA_FR"]["A2"].value, "CAFR-SKU")
            self.assertEqual(workbook["CA"].max_row, 1)
            workbook.close()
        finally:
            for path in (us_json, ca_fr_json, output):
                path.unlink(missing_ok=True)


if __name__ == "__main__":
    unittest.main()
