"""Validate CTA labels at the URLs in CTA_Verification.xlsx column K."""

import argparse
import html as html_lib
import json
import re
import shutil
import time
import unicodedata
from collections import Counter, defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timezone
from html.parser import HTMLParser
from pathlib import Path
from urllib.error import HTTPError, URLError
from urllib.parse import parse_qsl, urlencode, urlsplit, urlunsplit
from urllib.request import Request, urlopen

from openpyxl import load_workbook


USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
    "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0 Safari/537.36"
)
PD_SUFFIX = "(PD\ud310\uc815)"
SOURCE_SHEET = "CTA 판정결과"
RESULT_HEADERS = (
    "판정 모델코드", "모델 적용 URL", "CTA 표기 기대값", "실제 SKU", "실제 CTA",
    "CTA 활성", "CTA 검증", "검증 방식", "검증 상세", "확인 시각",
)


def text(value):
    return str(value or "").strip()


def normalize_country(value):
    return text(value).casefold().replace("-", "_")


def normalize_label(value):
    value = unicodedata.normalize("NFKC", text(value)).casefold()
    value = re.sub(r"[‐‑‒–—−-]+", " ", value)
    return " ".join(value.split())


CA_FR_LABEL_ALIASES = {
    "add to cart": "add_to_cart",
    "ajouter au panier": "add_to_cart",
    "where to buy": "where_to_buy",
    "où acheter": "where_to_buy",
    "ou acheter": "where_to_buy",
    # Some cached HTML is decoded with a replacement marker for the accent.
    "o? acheter": "where_to_buy",
    "notify me": "notify_me",
    "avisez moi": "notify_me",
    "m'aviser": "notify_me",
    "m’aviser": "notify_me",
    "m'avertir": "notify_me",
    "m’avertir": "notify_me",
    "avertissez moi": "notify_me",
    "me prévenir": "notify_me",
    "me prevenir": "notify_me",
    "buy now": "buy_now",
    "acheter maintenant": "buy_now",
    "achetez maintenant": "buy_now",
    "buy": "buy",
    "acheter": "buy",
    "achetez": "buy",
    "pre order": "pre_order",
    "pre order now": "pre_order",
    "pré commander": "pre_order",
    "précommander": "pre_order",
    "précommandez": "pre_order",
    "précommander maintenant": "pre_order",
    "précommandez maintenant": "pre_order",
}


def canonical_label(value, country=None):
    normalized = normalize_label(value)
    if normalize_country(country) == "ca_fr":
        return CA_FR_LABEL_ALIASES.get(normalized, normalized)
    return normalized


def labels_match(actual, expected, country=None):
    return canonical_label(actual, country) == canonical_label(expected, country)


def unverified_failure_verdict(used_pd_fallback):
    """Treat an unusable PD fallback as a failed verdict, not an unknown one."""
    return "FAIL" + PD_SUFFIX if used_pd_fallback else "NOT_FOUND"


def is_no_rule_expected(value):
    """Return true for the source value meaning that no CTA is expected."""
    normalized = text(value)
    # The workbook's Korean label may be mojibake in the current environment,
    # but it remains the only expected value wrapped in parentheses.
    return len(normalized) >= 2 and normalized.startswith("(") and normalized.endswith(")")


def normalize_sku(value):
    return text(value).upper().replace(" ", "")


def with_model_code(url, model_code):
    parts = urlsplit(url)
    query = dict(parse_qsl(parts.query, keep_blank_values=True))
    query["modelCode"] = model_code
    return urlunsplit(
        (parts.scheme, parts.netloc, parts.path, urlencode(query), parts.fragment)
    )


def cache_key(model_code, url):
    return f"{normalize_sku(model_code)}|{url}"


class ElementCollector(HTMLParser):
    """Collect button/link text while retaining attributes from server HTML."""

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.stack = []
        self.items = []

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        capture = tag in {"button", "a"} or attrs.get("role") == "button"
        item = {"tag": tag, "attrs": attrs, "parts": []} if capture else None
        self.stack.append(item)

    def handle_startendtag(self, tag, attrs):
        self.handle_starttag(tag, attrs)
        self.handle_endtag(tag)

    def handle_data(self, data):
        for item in self.stack:
            if item is not None:
                item["parts"].append(data)

    def handle_endtag(self, tag):
        if not self.stack:
            return
        item = self.stack.pop()
        if item is not None:
            item["text"] = " ".join("".join(item["parts"]).split())
            self.items.append(item)


CTA_WORDS = re.compile(
    r"add to cart|buy now|\bbuy\b|notify me|pre[\s-]?order|where to buy|"
    r"ajouter au panier|acheter|achetez|pré[\s-]?commander|où acheter|avisez-moi|"
    r"m['’]avertir|m['’]aviser|me pr[ée]venir|avertissez-moi",
    re.I,
)


def candidate_label(item):
    attrs = item["attrs"]
    return text(item.get("text") or attrs.get("aria-label") or attrs.get("title"))


def candidate_score(item, expected=None):
    attrs = item["attrs"]
    classes = text(attrs.get("class")).casefold()
    href = text(attrs.get("href")).casefold()
    label = candidate_label(item)
    normalized = normalize_label(label)
    if not label or len(label) > 100:
        return -1
    if "utility-cart" in classes or "global-cart" in classes:
        return -1
    if "frequently-bought-together" in classes or "recommendation" in classes:
        return -1
    if (
        "footer" in classes
        or "buy-direct-get-more" in href
        or "buy direct get more" in normalized
        or "achetez directement obtenez plus" in normalized
    ):
        return -1
    if "business" in href and "for business" in normalized:
        return -1
    if normalized in {"menu", "open my menu", "close my menu"}:
        return -1
    if label.endswith("?") or re.match(
        r"^(?:what|which|how|why|when|where can|can i|do i|is there)\b",
        normalized,
    ):
        return -1

    signal_score = 0
    signals = {
        "summary_ctaatcbutton": 240,
        "summary_cta": 220,
        "tg-add-to-cart": 230,
        "js-buy-now": 220,
        "tg-wtb": 230,
        "js-cta-buy": 220,
        "shop-status__buy-now": 210,
        "floating-navigation__button": 200,
        "watch-bc-buyflow": 190,
        "cta-wrapper": 180,
        "where-to-buy": 170,
        "anchorbtn": 190,
        "js-pfv2-buy-now": 180,
    }
    for signal, points in signals.items():
        if signal in classes:
            signal_score = max(signal_score, points)
    has_cta_text = bool(CTA_WORDS.search(label))
    if not signal_score and not has_cta_text:
        return -1

    score = signal_score
    if item["tag"] == "button":
        score += 25
    if attrs.get("role") == "button":
        score += 15
    if has_cta_text:
        score += 40
    # Disambiguate two primary product CTAs, but never promote an unrelated
    # page button merely because its text matches the expected answer.
    if (
        expected
        and signal_score
        and canonical_label(label, "ca_fr") == canonical_label(expected, "ca_fr")
    ):
        score += 120
    return score if score >= 40 else -1


def cache_item_usable(item):
    """Reject stale cache entries created from non-product menu/FAQ buttons."""
    if not item:
        return False
    actual = item.get("actual_cta")
    if not actual:
        return True
    candidate = item.get("candidate") or {}
    cached_item = {
        "tag": candidate.get("tag") or "button",
        "attrs": {
            "class": candidate.get("class"),
            "href": candidate.get("href"),
        },
        "text": actual,
    }
    return candidate_score(cached_item) >= 0


def needs_browser_confirmation(item, expected, country, mode):
    """Return whether an HTTP/cache result still needs rendered-DOM confirmation."""
    if mode == "none":
        return False
    if item.get("method") == "browser-dom" and cache_item_usable(item):
        return False
    unknown = not item.get("actual_cta")
    if unknown and is_no_rule_expected(expected):
        return False
    if mode == "unknown":
        return unknown
    return unknown or not labels_match(item.get("actual_cta"), expected, country)


def product_json_ld(html):
    scripts = re.findall(
        r"<script\b[^>]*type=[\"']application/ld\+json[\"'][^>]*>(.*?)</script>",
        html,
        flags=re.I | re.S,
    )
    for raw in scripts:
        try:
            value = json.loads(html_lib.unescape(raw))
        except (TypeError, ValueError):
            continue
        values = value if isinstance(value, list) else [value]
        for item in values:
            if isinstance(item, dict) and item.get("@type") == "Product":
                return item
    return {}


def extract_from_html(html, expected=None):
    parser = ElementCollector()
    parser.feed(html)
    ranked = sorted(
        ((candidate_score(item, expected), item) for item in parser.items),
        key=lambda pair: pair[0],
        reverse=True,
    )
    ranked = [pair for pair in ranked if pair[0] >= 0]
    product = product_json_ld(html)
    if not ranked:
        return {
            "actual_cta": None, "enabled": None, "actual_sku": product.get("sku"),
            "candidate": None,
        }
    score, item = ranked[0]
    attrs = item["attrs"]
    classes = text(attrs.get("class")).casefold()
    disabled = (
        "disabled" in attrs or text(attrs.get("aria-disabled")).casefold() == "true"
        or "disable" in classes
    )
    return {
        "actual_cta": candidate_label(item),
        "enabled": not disabled,
        "actual_sku": product.get("sku"),
        "candidate": {
            "tag": item["tag"], "class": attrs.get("class"),
            "href": attrs.get("href"), "score": score,
        },
    }


def fetch_url(url, expected, timeout, retries=2):
    error = None
    status = None
    for attempt in range(retries + 1):
        try:
            request = Request(
                url, headers={"User-Agent": USER_AGENT, "Accept-Language": "en-US,en;q=0.8"}
            )
            with urlopen(request, timeout=timeout) as response:
                status = response.status
                html = response.read(8_000_000).decode("utf-8", "replace")
                result = extract_from_html(html, expected)
                result.update({
                    "method": "http-html", "http_status": status,
                    "final_url": response.geturl(), "error": None,
                    "checked_at": datetime.now(timezone.utc).isoformat(),
                })
                return result
        except (HTTPError, URLError, TimeoutError, OSError) as exc:
            status = getattr(exc, "code", None)
            error = f"{type(exc).__name__}: {exc}"
            if attempt < retries:
                time.sleep(0.5 * (attempt + 1))
    return {
        "actual_cta": None, "enabled": None, "actual_sku": None,
        "candidate": None, "method": "http-html", "http_status": status,
        "final_url": url, "error": error,
        "checked_at": datetime.now(timezone.utc).isoformat(),
    }


def browser_extract(page, url, expected, timeout_ms):
    response = page.goto(url, wait_until="domcontentloaded", timeout=timeout_ms)
    page.wait_for_timeout(3500)
    values = page.evaluate(
        """(expected) => {
        const visible = e => {
          const r=e.getBoundingClientRect(), s=getComputedStyle(e);
          return r.width>0 && r.height>0 && s.display!=='none' && s.visibility!=='hidden';
        };
        const nodes=[...document.querySelectorAll('button,a,[role="button"]')]
          .filter(visible).map(e=>({tag:e.tagName.toLowerCase(),text:(e.innerText||e.getAttribute('aria-label')||e.title||'').trim().replace(/\\s+/g,' '),attrs:{class:typeof e.className==='string'?e.className:'',href:e.getAttribute('href'),role:e.getAttribute('role'),'aria-disabled':e.getAttribute('aria-disabled'),disabled:e.disabled?'':null}}));
        const scripts=[...document.querySelectorAll('script[type="application/ld+json"]')];
        let sku=null;
        for(const s of scripts){try{let x=JSON.parse(s.textContent);for(const y of(Array.isArray(x)?x:[x]))if(y?.['@type']==='Product')sku=y.sku||sku}catch(e){}}
        return {nodes,sku};
        }""",
        expected,
    )
    parser_items = [
        {"tag": node["tag"], "attrs": node["attrs"], "text": node["text"], "parts": []}
        for node in values["nodes"]
    ]
    ranked = sorted(
        ((candidate_score(item, expected), item) for item in parser_items),
        key=lambda pair: pair[0], reverse=True,
    )
    ranked = [pair for pair in ranked if pair[0] >= 0]
    if ranked:
        score, item = ranked[0]
        attrs = item["attrs"]
        classes = text(attrs.get("class")).casefold()
        disabled = (
            attrs.get("disabled") is not None
            or text(attrs.get("aria-disabled")).casefold() == "true"
            or "disable" in classes
        )
        actual, enabled = candidate_label(item), not disabled
        candidate = {"tag": item["tag"], "class": attrs.get("class"), "href": attrs.get("href"), "score": score}
    else:
        actual, enabled, candidate = None, None, None
    return {
        "actual_cta": actual, "enabled": enabled, "actual_sku": values["sku"],
        "candidate": candidate, "method": "browser-dom",
        "http_status": response.status if response else None, "final_url": page.url,
        "error": None, "checked_at": datetime.now(timezone.utc).isoformat(),
    }


def load_cache(path):
    try:
        return json.loads(path.read_text(encoding="utf-8")) if path.exists() else {}
    except (OSError, ValueError):
        return {}


def save_cache(path, cache):
    path.write_text(json.dumps(cache, ensure_ascii=False, indent=2), encoding="utf-8")


def _pf_country(record):
    source = text(record.get("source_pf_url") or record.get("category_url"))
    match = re.search(r"/(us|ca_fr|ca)(?:/|$)", source, re.I)
    return normalize_country(match.group(1)) if match else ""


def load_pf_data(paths):
    """Load fresh PF-card CTA records, grouped by country and exact SKU."""
    grouped = defaultdict(list)
    ignored_legacy = 0
    for raw_path in paths:
        path = Path(raw_path).resolve()
        data = json.loads(path.read_text(encoding="utf-8-sig"))
        records = data.get("records", []) if isinstance(data, dict) else data
        for record in records:
            if not isinstance(record, dict) or record.get("status"):
                continue
            if "pf_cta_label" not in record:
                ignored_legacy += 1
                continue
            country = _pf_country(record)
            sku = text(record.get("sku") or record.get("model_code") or record.get("modelcode"))
            if country and sku:
                grouped[(country, normalize_sku(sku))].append(record)

    resolved = {}
    conflicts = 0
    for key, records in grouped.items():
        country = key[0]
        semantic_counts = Counter(
            canonical_label(record.get("pf_cta_label"), country) for record in records
        )
        chosen_semantic, _count = semantic_counts.most_common(1)[0]
        candidates = [
            record for record in records
            if canonical_label(record.get("pf_cta_label"), country) == chosen_semantic
        ]
        candidates.sort(
            key=lambda record: (
                (record.get("sort_type") or "").casefold() != "recommended",
                not bool(record.get("is_default")),
            )
        )
        chosen = candidates[0]
        distinct = sorted({semantic for semantic in semantic_counts if semantic})
        if len(distinct) > 1:
            conflicts += 1
        source = text(chosen.get("source_pf_url"))
        if source.startswith("/"):
            source = "https://www.samsung.com" + source.rstrip("/") + "/"
        resolved[key] = {
            "actual_cta": chosen.get("pf_cta_label"),
            "enabled": chosen.get("pf_cta_enabled"),
            "actual_sku": chosen.get("sku") or chosen.get("model_code"),
            "candidate": {
                "class": chosen.get("pf_cta_class"),
                "source_pf_url": chosen.get("source_pf_url"),
                "sort_type": chosen.get("sort_type"),
            },
            "method": "pf-card-dom",
            "http_status": None,
            "final_url": source,
            "error": None,
            "checked_at": datetime.now(timezone.utc).isoformat(),
            "pf_semantic_conflicts": distinct,
        }
    return resolved, ignored_legacy, conflicts


def choose_stratified(rows, count):
    groups = defaultdict(list)
    for row in rows:
        groups[normalize_label(row["expected"]) or "NO_RULE"].append(row)
    chosen = []
    while len(chosen) < count and groups:
        for key in list(groups):
            if groups[key] and len(chosen) < count:
                chosen.append(groups[key].pop(0))
            if not groups[key]:
                del groups[key]
    return chosen


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("input_xlsx")
    ap.add_argument("--output")
    ap.add_argument("--countries", default="us,ca,ca_fr")
    ap.add_argument("--workers", type=int, default=6)
    ap.add_argument("--timeout", type=float, default=30)
    ap.add_argument("--log-every", type=int, default=25)
    ap.add_argument("--cache", default="cta_label_cache.json")
    ap.add_argument(
        "--pf-data", action="append", default=[],
        help="fresh PF scrape JSON; repeat for multiple countries",
    )
    ap.add_argument("--refresh", action="store_true")
    ap.add_argument("--sample", type=int, help="stratified sample size")
    ap.add_argument("--exclude-expected", action="append", default=[])
    ap.add_argument(
        "--browser-confirm", choices=("none", "unknown", "mismatch"), default="mismatch",
        help="render unknown pages only, or unknown and HTTP mismatches",
    )
    args = ap.parse_args()
    args.log_every = max(1, args.log_every)

    input_path = Path(args.input_xlsx).resolve()
    output_path = Path(args.output).resolve() if args.output else input_path.with_name(input_path.stem + "_checked.xlsx")
    cache_path = Path(args.cache).resolve()
    countries = {normalize_country(x) for x in args.countries.split(",") if text(x)}
    print(
        f"[start] input={input_path.name} countries={','.join(sorted(countries))} "
        f"workers={max(1, args.workers)} browser_confirm={args.browser_confirm}",
        flush=True,
    )
    print(f"[start] output={output_path} cache={cache_path}", flush=True)

    wb = load_workbook(input_path, data_only=False)
    ws = wb[SOURCE_SHEET]
    header = [cell.value for cell in ws[1]]
    index = {name: number + 1 for number, name in enumerate(header)}
    required = ["국가", "모델코드", "CTA 표기", "URL"]
    missing = [name for name in required if name not in index]
    if missing:
        raise SystemExit(f"missing columns: {missing}")

    rows = []
    excluded = {normalize_label(x) for x in args.exclude_expected}
    for number in range(2, ws.max_row + 1):
        country = normalize_country(ws.cell(number, index["국가"]).value)
        if country not in countries:
            continue
        code = text(ws.cell(number, index["모델코드"]).value).upper()
        if not code:
            continue
        url = text(ws.cell(number, index["URL"]).value)
        expected = text(ws.cell(number, index["CTA 표기"]).value) or None
        if expected == "(규칙 미해당)":
            expected = None
        if normalize_label(expected) in excluded:
            continue
        rows.append({
            "row": number, "country": country, "code": code, "url": url,
            "expected": expected,
        })
    if args.sample:
        rows = choose_stratified(rows, args.sample)
    unique_model_urls = {(r["code"], r["url"]) for r in rows if r["url"]}
    country_counts = Counter(row["country"] for row in rows)
    expected_counts = Counter(row["expected"] or "NO_RULE" for row in rows)
    print(
        f"[scope] rows={len(rows)} unique_model_urls={len(unique_model_urls)} "
        f"countries={dict(country_counts)}",
        flush=True,
    )
    print(f"[scope] expected_labels={dict(expected_counts)}", flush=True)

    pf_items, ignored_legacy_pf, pf_conflicts = load_pf_data(args.pf_data)
    pf_scope = {
        key
        for key in ((row["country"], normalize_sku(row["code"])) for row in rows)
        if key in pf_items
    }
    if args.pf_data:
        fallback_rows = sum(
            (row["country"], normalize_sku(row["code"])) not in pf_items
            for row in rows
        )
        print(
            f"[pf] matched_skus={len(pf_scope)} fallback_rows={fallback_rows} "
            f"legacy_records_ignored={ignored_legacy_pf} conflicts={pf_conflicts}",
            flush=True,
        )

    cache = load_cache(cache_path)
    work = {}
    for row in rows:
        pf_key = (row["country"], normalize_sku(row["code"]))
        if pf_key not in pf_items and row["url"]:
            work.setdefault(
                (row["code"], row["url"]),
                {"expected": row["expected"], "country": row["country"]},
            )
    pending = [
        (code, url, value["expected"])
        for (code, url), value in work.items()
        if args.refresh or not cache_item_usable(cache.get(cache_key(code, url)))
    ]
    invalid_cached = sum(
        1 for code, url in work
        if cache_key(code, url) in cache
        and not cache_item_usable(cache.get(cache_key(code, url)))
    )
    print(
        f"[http] start total={len(work)} cached={len(work) - len(pending)} "
        f"pending={len(pending)} invalid_cached={invalid_cached} "
        f"refresh={args.refresh}",
        flush=True,
    )
    started = time.time()
    http_counts = Counter()
    with ThreadPoolExecutor(max_workers=max(1, args.workers)) as executor:
        futures = {
            executor.submit(
                fetch_url, with_model_code(url, code), expected, args.timeout
            ): (code, url)
            for code, url, expected in pending
        }
        for done, future in enumerate(as_completed(futures), 1):
            code, url = futures[future]
            result = future.result()
            cache[cache_key(code, url)] = result
            http_counts["cta_found" if result.get("actual_cta") else "cta_missing"] += 1
            if result.get("error"):
                http_counts["error"] += 1
            if done % args.log_every == 0 or done == len(pending):
                elapsed = max(time.time() - started, 0.01)
                eta = (len(pending) - done) / (done / elapsed) if done else 0
                print(
                    f"[http] {done}/{len(pending)} elapsed={elapsed/60:.1f}m "
                    f"eta={eta/60:.1f}m last_model={code} stats={dict(http_counts)}",
                    flush=True,
                )
                save_cache(cache_path, cache)

    browser_urls = []
    for (code, url), value in work.items():
        expected = value["expected"]
        country = value["country"]
        item = cache[cache_key(code, url)]
        if needs_browser_confirmation(
            item, expected, country, args.browser_confirm
        ):
            browser_urls.append((code, url, expected))
    if browser_urls:
        from playwright.sync_api import sync_playwright
        print(f"[browser] start confirm={len(browser_urls)}", flush=True)
        browser_started = time.time()
        browser_counts = Counter()
        with sync_playwright() as playwright:
            browser = playwright.chromium.launch(channel="chrome", headless=True)
            page = browser.new_page(user_agent=USER_AGENT, viewport={"width": 1440, "height": 1000})
            for done, (code, url, expected) in enumerate(browser_urls, 1):
                key = cache_key(code, url)
                try:
                    cache[key] = browser_extract(
                        page, with_model_code(url, code), expected,
                        int(args.timeout * 1000),
                    )
                except Exception as exc:  # noqa: BLE001
                    cache[key] = {
                        "actual_cta": None, "enabled": None, "actual_sku": None,
                        "candidate": None, "method": "browser-dom", "http_status": None,
                        "final_url": page.url, "error": f"{type(exc).__name__}: {exc}",
                        "checked_at": datetime.now(timezone.utc).isoformat(),
                    }
                result = cache[key]
                browser_counts["cta_found" if result.get("actual_cta") else "cta_missing"] += 1
                if result.get("error"):
                    browser_counts["error"] += 1
                if done % min(args.log_every, 10) == 0 or done == len(browser_urls):
                    elapsed = max(time.time() - browser_started, 0.01)
                    eta = (len(browser_urls) - done) / (done / elapsed)
                    print(
                        f"[browser] {done}/{len(browser_urls)} elapsed={elapsed/60:.1f}m "
                        f"eta={eta/60:.1f}m last_model={code} stats={dict(browser_counts)}",
                        flush=True,
                    )
                    save_cache(cache_path, cache)
            browser.close()
    else:
        print("[browser] confirm=0; browser phase skipped", flush=True)
    save_cache(cache_path, cache)

    print(f"[write] writing {len(rows)} judged rows", flush=True)
    start_col = index.get(RESULT_HEADERS[0], ws.max_column + 1)
    for offset, name in enumerate(RESULT_HEADERS):
        ws.cell(1, start_col + offset).value = name
    counts = Counter()
    for row in rows:
        applied_url = with_model_code(row["url"], row["code"]) if row["url"] else None
        pf_key = (row["country"], normalize_sku(row["code"]))
        pf_item = pf_items.get(pf_key)
        used_pd_fallback = bool(args.pf_data) and pf_item is None
        item = pf_item or (cache.get(cache_key(row["code"], row["url"]), {}) if row["url"] else {})
        actual = item.get("actual_cta")
        expected = row["expected"]
        actual_sku = item.get("actual_sku")
        if used_pd_fallback and not row["url"]:
            verdict = "NOT_FOUND"
            reason = "SKU not found on PF and column K URL is empty"
        elif not args.pf_data and not row["url"]:
            verdict = "NOT_FOUND"
            reason = "column K URL is empty"
        elif actual_sku and normalize_sku(actual_sku) != normalize_sku(row["code"]):
            verdict = unverified_failure_verdict(used_pd_fallback)
            reason = f"page SKU differs from column B: page={actual_sku}, B={row['code']}"
        elif item.get("error") and not actual:
            verdict = unverified_failure_verdict(used_pd_fallback)
            reason = item["error"]
        elif is_no_rule_expected(expected) and not actual:
            verdict = "NO_RULE"
            reason = (
                f"no CTA expected; http={item.get('http_status')}; "
                f"final={item.get('final_url')}"
            )
        else:
            label_match = labels_match(actual, expected, row["country"])
            enabled = item.get("enabled")
            verdict = "PASS" if label_match else "FAIL"
            if used_pd_fallback:
                verdict += PD_SUFFIX
            reason = (
                f"judgment_source={'PD fallback' if used_pd_fallback else 'PF card'}; "
                f"expected source=CTA 표기; label_match={label_match}; enabled={enabled}; "
                f"expected_semantic={canonical_label(expected, row['country'])}; "
                f"actual_semantic={canonical_label(actual, row['country'])}; "
                f"http={item.get('http_status')}; final={item.get('final_url')}; "
                f"candidate={item.get('candidate')}; "
                f"pf_semantic_conflicts={item.get('pf_semantic_conflicts', [])}"
            )
        counts[verdict] += 1
        values = (
            row["code"], applied_url, expected, actual_sku, actual,
            item.get("enabled"), verdict,
            item.get("method"), reason, item.get("checked_at"),
        )
        for offset, value in enumerate(values):
            ws.cell(row["row"], start_col + offset).value = value

    if output_path == input_path:
        backup = input_path.with_name(input_path.stem + ".pre_cta_check.xlsx")
        if not backup.exists():
            shutil.copy2(input_path, backup)
    wb.save(output_path)
    print(f"[done] results={dict(counts)}", flush=True)
    print(f"[done] output={output_path}", flush=True)
    print(f"[done] cache={cache_path}", flush=True)


if __name__ == "__main__":
    main()
