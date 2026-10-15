import json
from pathlib import Path
import unittest
from uuid import uuid4

from flask import Flask
from openpyxl import Workbook

import config
from product_catalog import connection, import_country, product_detail, search_catalog
from routes.products import products_bp


class ProductCatalogTests(unittest.TestCase):
    def setUp(self):
        stem = Path(__file__).parent / ("_catalog_test_" + uuid4().hex)
        self.db = stem.with_suffix(".db")
        self.workbook = stem.with_suffix(".xlsx")
        self.long_file = stem.with_suffix(".jsonl")
        self.paths = [self.db, self.workbook, self.long_file, Path(str(self.db)+"-wal"), Path(str(self.db)+"-shm")]
        for path in self.paths:
            self.addCleanup(path.unlink, missing_ok=True)
        self.write_source()
        import_country(self.db,"us",self.workbook,self.long_file,log=lambda *a,**k: None)

    def write_source(self, bad_column=False):
        wb = Workbook()
        sheet = wb.active
        sheet.title = "sr_merged_product"
        sheet.append(["model_code","model_name","product_name","product_color","product_specification","_ctid"])
        sheet.append(["SKU/AA","BASE-1","Phone sample","Violet Shadow",json.dumps({"Display":{"Resolution":"2504 x 2256"},"Storage":{"Memory":"12 GB"}}),"ctid:(0,1)"])
        sheet = wb.create_sheet("product_spec_basics")
        sheet.append(["model_code","display_name","feature_json","_ctid"])
        sheet.append(["SKU/AA","Other name",self.long_file.name+":L1","ctid:(0,1)"])
        sheet = wb.create_sheet("sr_product_description")
        sheet.append(["model_code","content","_ctid"])
        sheet.append(["SKU/AA",self.long_file.name+":L2","ctid:(0,1)"])
        wb.save(self.workbook)
        wb.close()
        feature = {"FeatureItem":[{"FeatureTitle":{"featureTitle":"<b>한국어 카메라</b>"},"FeatureBody":{"featureBody":"Ultra camera detail"},"FeatureMedia":[{"url":"https://example.invalid/image.png"}]}]}
        values = [
            {"table":"product_spec_basics","column":"wrong" if bad_column else "feature_json","ctid":"ctid:(0,1)","value":json.dumps(feature)},
            {"table":"sr_product_description","column":"content","ctid":"ctid:(0,1)","value":"## Original product description\nOffline reference."},
        ]
        self.long_file.write_text("\n".join(json.dumps(r) for r in values),encoding="utf-8")

    def test_parsed_specs_and_provenance(self):
        detail = product_detail(self.db,"us","SKU/AA","specs")
        self.assertEqual(detail["sources"][0]["specs"][0]["value"],"2504 x 2256")
        self.assertEqual(detail["sources"][0]["row_number"],2)
        self.assertEqual(detail["sources"][0]["field_name"],"product_specification")

    def test_long_reference_and_local_media_only(self):
        source = product_detail(self.db,"us","SKU/AA","features")["sources"][0]
        self.assertEqual(source["long_line"],1)
        self.assertEqual(source["features"][0]["title"],"한국어 카메라")
        self.assertEqual(source["features"][0]["media"][0]["url"],"https://example.invalid/image.png")

    def test_indexed_search_scopes_and_literal_input(self):
        self.assertEqual(search_catalog(self.db,query="SKU/AA")["total"],1)
        self.assertEqual(search_catalog(self.db,query="Resolution",scope="specs")["total"],1)
        self.assertEqual(search_catalog(self.db,query="카메라",scope="features")["total"],1)
        self.assertEqual(search_catalog(self.db,query="offline",scope="description")["total"],1)
        self.assertEqual(search_catalog(self.db,query="' OR 1=1 --")["total"],0)
        self.assertEqual(search_catalog(self.db,query="%_")["total"],0)

    def test_repeat_import_replaces_country_without_duplicate_documents(self):
        import_country(self.db,"ca",self.workbook,self.long_file,log=lambda *a,**k:None)
        import_country(self.db,"us",self.workbook,self.long_file,log=lambda *a,**k:None)
        self.assertEqual(search_catalog(self.db)["total"],2)
        with connection(self.db) as conn:
            self.assertEqual(conn.execute("SELECT COUNT(*) FROM documents").fetchone()[0],3)

    def test_failed_import_keeps_previous_country_data(self):
        self.write_source(bad_column=True)
        with self.assertRaises(ValueError):
            import_country(self.db,"us",self.workbook,self.long_file,log=lambda *a,**k:None)
        self.assertEqual(search_catalog(self.db,query="SKU/AA")["total"],1)
        self.assertEqual(len(product_detail(self.db,"us","SKU/AA","features")["sources"]),1)

    def test_changed_source_drops_orphaned_text_and_search_entries(self):
        contents = self.long_file.read_text(encoding="utf-8").replace("Ultra camera detail", "Updated camera detail")
        self.long_file.write_text(contents,encoding="utf-8")
        import_country(self.db,"us",self.workbook,self.long_file,log=lambda *a,**k:None)
        self.assertEqual(search_catalog(self.db,query="Ultra",scope="features")["total"],0)
        self.assertEqual(search_catalog(self.db,query="Updated",scope="features")["total"],1)
        with connection(self.db) as conn:
            self.assertEqual(conn.execute("SELECT COUNT(*) FROM documents").fetchone()[0],3)

    def test_api_login_validation_and_slash_sku(self):
        app = Flask(__name__)
        app.secret_key = "test-only"
        app.register_blueprint(products_bp)
        original = config.PRODUCT_DB_PATH
        config.PRODUCT_DB_PATH = str(self.db)
        self.addCleanup(setattr, config, "PRODUCT_DB_PATH", original)
        client = app.test_client()
        self.assertEqual(client.get("/api/products").status_code,401)
        with client.session_transaction() as session:
            session["nickname"] = "tester"
        self.assertEqual(client.get("/api/products?country=bad").status_code,400)
        self.assertEqual(client.get("/api/products?page=oops").status_code,400)
        self.assertEqual(client.get("/api/products/us/SKU%2FAA?section=features").status_code,200)
        self.assertEqual(client.get("/api/products/us/MISSING").status_code,404)
        exported = client.get("/api/products/export?q=SKU%2FAA")
        self.assertEqual(exported.status_code,200)
        self.assertTrue(exported.data.startswith(b"PK"))


if __name__ == "__main__":
    unittest.main()
