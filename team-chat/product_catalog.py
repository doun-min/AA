"""Local product catalog, source provenance and indexed long-text search."""

from contextlib import contextmanager
from datetime import datetime, timezone
import hashlib
from html import unescape
from html.parser import HTMLParser
import json
from pathlib import Path
import re
import sqlite3


COUNTRIES = ("us", "ca", "ca_fr")
SECTIONS = ("basic", "specs", "features", "description")
TABLES = (
    "sr_merged_product", "product_spec_basics", "product_family_list",
    "product_spec_color", "model_price", "product_comment_statistics",
    "sr_product_description",
)
SCHEMA = """
PRAGMA foreign_keys=ON;
CREATE TABLE IF NOT EXISTS products (
    country TEXT NOT NULL, sku TEXT NOT NULL, name TEXT NOT NULL DEFAULT '',
    model_name TEXT NOT NULL DEFAULT '', color TEXT NOT NULL DEFAULT '',
    category TEXT NOT NULL DEFAULT '', PRIMARY KEY(country,sku)
);
CREATE TABLE IF NOT EXISTS source_records (
    id INTEGER PRIMARY KEY, country TEXT NOT NULL, sku TEXT NOT NULL,
    source_file TEXT NOT NULL, table_name TEXT NOT NULL, row_number INTEGER NOT NULL,
    ctid TEXT, payload TEXT NOT NULL,
    FOREIGN KEY(country,sku) REFERENCES products(country,sku) ON DELETE CASCADE
);
CREATE INDEX IF NOT EXISTS idx_source_sku ON source_records(country,sku);
CREATE TABLE IF NOT EXISTS documents (
    id INTEGER PRIMARY KEY, section TEXT NOT NULL, digest TEXT NOT NULL,
    raw TEXT NOT NULL, plain TEXT NOT NULL, UNIQUE(section,digest)
);
CREATE VIRTUAL TABLE IF NOT EXISTS document_search USING fts5(
    plain, tokenize='unicode61 remove_diacritics 2'
);
CREATE TABLE IF NOT EXISTS product_documents (
    source_id INTEGER NOT NULL REFERENCES source_records(id) ON DELETE CASCADE,
    document_id INTEGER NOT NULL REFERENCES documents(id),
    field_name TEXT NOT NULL, long_file TEXT, long_line INTEGER,
    PRIMARY KEY(source_id,field_name)
);
CREATE INDEX IF NOT EXISTS idx_document_links ON product_documents(document_id);
CREATE TABLE IF NOT EXISTS spec_values (
    document_id INTEGER NOT NULL REFERENCES documents(id),
    spec_group TEXT NOT NULL, spec_name TEXT NOT NULL, spec_value TEXT NOT NULL
);
CREATE INDEX IF NOT EXISTS idx_spec_document ON spec_values(document_id);
CREATE TABLE IF NOT EXISTS imports (
    country TEXT PRIMARY KEY, workbook TEXT NOT NULL, long_file TEXT,
    imported_at TEXT NOT NULL, product_count INTEGER NOT NULL,
    source_rows INTEGER NOT NULL, unresolved_refs INTEGER NOT NULL
);
"""


class PlainHTML(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts = []
        self.hidden = 0

    def handle_starttag(self, tag, attrs):
        if tag in {"script", "style"}:
            self.hidden += 1
        if tag in {"p", "br", "div", "li"}:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in {"script", "style"}:
            self.hidden = max(0, self.hidden - 1)

    def handle_data(self, data):
        if not self.hidden:
            self.parts.append(data)


def clean_text(value):
    parser = PlainHTML()
    parser.feed(str(value or ""))
    return unescape("".join(parser.parts)).replace("\xa0", " ").strip()


def json_value(value):
    if isinstance(value, str):
        try:
            return json.loads(value)
        except (ValueError, TypeError):
            pass
    return value


def display_value(value):
    return json.dumps(value, ensure_ascii=False, default=str) if isinstance(value, (dict, list)) else str(value if value is not None else "")


def flatten_specs(value, path=()):
    value = json_value(value)
    if isinstance(value, dict):
        for key, child in value.items():
            yield from flatten_specs(child, path + (str(key),))
    elif isinstance(value, list):
        for index, child in enumerate(value):
            yield from flatten_specs(child, path + (str(index + 1),))
    elif value is not None:
        yield {
            "group": " / ".join(path[:-1]),
            "name": path[-1] if path else "Specification",
            "value": display_value(value),
        }


def feature_items(value):
    """Extract text/media references only; never fetch or execute source HTML."""
    output = []

    def walk(node):
        if isinstance(node, list):
            for item in node:
                walk(item)
        elif isinstance(node, dict):
            if "FeatureTitle" in node or "FeatureBody" in node:
                def text_from(key, child_key):
                    raw = node.get(key)
                    return clean_text(raw.get(child_key) if isinstance(raw, dict) else raw)
                media = node.get("FeatureMedia") or []
                if isinstance(media, dict):
                    media = [media]
                item = {
                    "title": text_from("FeatureTitle", "featureTitle"),
                    "subtitle": text_from("FeatureSubTitle", "featureSubTitle"),
                    "body": text_from("FeatureBody", "featureBody"),
                    "footnote": clean_text(node.get("FeatureFootnote") or node.get("FeatureDisclaimer")),
                    "media": [
                        {"url": entry.get("url"), "description": clean_text(entry.get("description") or entry.get("substitueImgAlt"))}
                        for entry in media if isinstance(entry, dict) and entry.get("url")
                    ],
                }
                if any(item.values()):
                    output.append(item)
            for key, child in node.items():
                if key not in {"FeatureTitle", "FeatureSubTitle", "FeatureBody", "FeatureMedia"}:
                    walk(child)

    walk(json_value(value))
    return output


def searchable_text(section, value):
    parsed = json_value(value)
    if section == "specs":
        return "\n".join(f"{r['group']} {r['name']} {r['value']}" for r in flatten_specs(parsed))
    if section == "features":
        return "\n".join(" ".join([r["title"], r["subtitle"], r["body"], r["footnote"]]) for r in feature_items(parsed))
    return clean_text(display_value(parsed))


@contextmanager
def connection(path, write=False):
    path = Path(path)
    if write:
        path.parent.mkdir(parents=True, exist_ok=True)
        conn = sqlite3.connect(path, timeout=30)
        conn.execute("PRAGMA journal_mode=WAL")
    else:
        conn = sqlite3.connect(path.resolve().as_uri() + "?mode=ro", uri=True, timeout=10)
        conn.execute("PRAGMA query_only=ON")
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys=ON")
    try:
        yield conn
    finally:
        conn.close()


def init_catalog(conn):
    conn.executescript(SCHEMA)


def add_document(conn, section, value):
    raw = display_value(value)
    digest = hashlib.sha256(raw.encode("utf-8")).hexdigest()
    existing = conn.execute("SELECT id FROM documents WHERE section=? AND digest=?", (section, digest)).fetchone()
    if existing:
        return existing["id"]
    plain = searchable_text(section, value)
    cur = conn.execute("INSERT INTO documents(section,digest,raw,plain) VALUES(?,?,?,?)", (section, digest, raw, plain))
    doc_id = cur.lastrowid
    conn.execute("INSERT INTO document_search(rowid,plain) VALUES(?,?)", (doc_id, plain))
    if section == "specs":
        conn.executemany("INSERT INTO spec_values VALUES(?,?,?,?)", [
            (doc_id, item["group"], item["name"], item["value"]) for item in flatten_specs(value)
        ])
    return doc_id


def like_value(value):
    return "%" + value.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_") + "%"


def search_catalog(path, country="", query="", scope="basic", page=1, page_size=50):
    if not Path(path).is_file():
        return {"items": [], "total": 0, "page": page, "page_size": page_size}
    clauses, args = [], []
    if country:
        clauses.append("p.country=?")
        args.append(country)
    if query:
        basic = "(p.sku LIKE ? ESCAPE '\\' OR p.name LIKE ? ESCAPE '\\' OR p.model_name LIKE ? ESCAPE '\\' OR p.color LIKE ? ESCAPE '\\')"
        basic_args = [like_value(query)] * 4
        tokens = re.findall(r"\w+", query, re.UNICODE)
        if scope == "basic" or not tokens:
            clauses.append(basic)
            args.extend(basic_args)
        else:
            match = " AND ".join('"' + word + '"*' for word in tokens)
            doc_where = ""
            doc_args = [match]
            if scope in {"specs", "features", "description"}:
                doc_where = " AND d.section=?"
                doc_args.append(scope)
            doc_search = """EXISTS (
                SELECT 1 FROM source_records r JOIN product_documents l ON l.source_id=r.id
                JOIN documents d ON d.id=l.document_id
                WHERE r.country=p.country AND r.sku=p.sku
                AND d.id IN (SELECT rowid FROM document_search WHERE document_search MATCH ?)
            """ + doc_where + ")"
            clauses.append(f"({basic} OR {doc_search})" if scope == "all" else doc_search)
            args.extend((basic_args if scope == "all" else []) + doc_args)
    where = " WHERE " + " AND ".join(clauses) if clauses else ""
    with connection(path) as conn:
        total = conn.execute("SELECT COUNT(*) FROM products p" + where, args).fetchone()[0]
        rows = conn.execute(
            "SELECT p.* FROM products p" + where + " ORDER BY p.country,p.sku LIMIT ? OFFSET ?",
            args + [page_size, (page-1)*page_size],
        ).fetchall()
    return {"items": [dict(row) for row in rows], "total": total, "page": page, "page_size": page_size}


def catalog_stats(path):
    if not Path(path).is_file():
        return {"countries": [], "total_products": 0, "size_bytes": 0}
    with connection(path) as conn:
        rows = [dict(row) for row in conn.execute("SELECT * FROM imports ORDER BY country")]
    return {"countries": rows, "total_products": sum(r["product_count"] for r in rows), "size_bytes": Path(path).stat().st_size}


def product_detail(path, country, sku, section="basic"):
    if not Path(path).is_file():
        return None
    with connection(path) as conn:
        row = conn.execute("SELECT * FROM products WHERE country=? AND sku=?", (country,sku)).fetchone()
        if not row:
            return None
        result = {"product": dict(row), "section": section, "sources": []}
        sources = conn.execute("SELECT * FROM source_records WHERE country=? AND sku=? ORDER BY id", (country,sku)).fetchall()
        if section == "basic":
            result["sources"] = [{**dict(r), "payload": json.loads(r["payload"])} for r in sources]
        else:
            links = conn.execute("""SELECT d.id,d.raw,d.plain,r.source_file,r.table_name,r.row_number,r.ctid,
                l.field_name,l.long_file,l.long_line FROM source_records r
                JOIN product_documents l ON l.source_id=r.id JOIN documents d ON d.id=l.document_id
                WHERE r.country=? AND r.sku=? AND d.section=? ORDER BY r.id,l.field_name""", (country,sku,section)).fetchall()
            for link in links:
                item = dict(link)
                if section == "specs":
                    item["specs"] = [dict(r) for r in conn.execute("SELECT spec_group AS 'group',spec_name AS name,spec_value AS value FROM spec_values WHERE document_id=?", (item["id"],))]
                elif section == "features":
                    item["features"] = feature_items(item["raw"])
                result["sources"].append(item)
    return result


def import_country(path, country, workbook_path, long_path=None, log=print):
    from openpyxl import load_workbook

    workbook_path = Path(workbook_path).resolve()
    long_path = Path(long_path).resolve() if long_path else None
    if country not in COUNTRIES:
        raise ValueError("Unsupported country")
    if not workbook_path.is_file() or (long_path and not long_path.is_file()):
        raise FileNotFoundError("Source file not found")
    unresolved, source_rows = 0, 0
    ref_pattern = re.compile(r"^(.+\.jsonl):L(\d+)$", re.I)
    pending = {}
    ctid_links = {}
    with connection(path, write=True) as conn:
        init_catalog(conn)
        conn.execute("BEGIN IMMEDIATE")
        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
        try:
            conn.execute("DELETE FROM products WHERE country=?", (country,))
            for table in TABLES:
                if table not in workbook.sheetnames:
                    continue
                iterator = workbook[table].iter_rows(values_only=True)
                headers = [str(v or "") for v in next(iterator)]
                if "model_code" not in headers:
                    continue
                count = 0
                for row_number, values in enumerate(iterator, 2):
                    payload = dict(zip(headers, values))
                    sku = str(payload.get("model_code") or "").strip().upper()
                    if not sku:
                        continue
                    conn.execute("INSERT OR IGNORE INTO products(country,sku) VALUES(?,?)", (country,sku))
                    name = payload.get("product_name") or payload.get("display_name")
                    updates = {
                        "name": name,
                        "model_name": payload.get("model_name") if table != "product_comment_statistics" else None,
                        "color": payload.get("product_color") or payload.get("colors"),
                        "category": payload.get("category_name") or payload.get("display_category_minor"),
                    }
                    for field, value in updates.items():
                        if value is not None and str(value).strip():
                            conn.execute(f"UPDATE products SET {field}=? WHERE country=? AND sku=? AND {field}=''", (str(value),country,sku))
                    documents = []
                    for field, value in list(payload.items()):
                        if value is None:
                            continue
                        section = {"product_specification": "specs", "feature_json": "features", "content": "description", "review_text_collection": "description"}.get(field)
                        if not section:
                            continue
                        payload[field] = {"section": section, "value_source": "detail"}
                        documents.append((field,section,value))
                    cur = conn.execute("INSERT INTO source_records(country,sku,source_file,table_name,row_number,ctid,payload) VALUES(?,?,?,?,?,?,?)", (country,sku,str(workbook_path),table,row_number,str(payload.get("_ctid") or ""),json.dumps(payload,ensure_ascii=False,default=str)))
                    source_id = cur.lastrowid
                    if table == "sr_product_description":
                        ctid_links.setdefault(str(payload.get("_ctid") or ""), []).append(source_id)
                    for field,section,value in documents:
                        ref = ref_pattern.match(str(value))
                        if ref:
                            pending.setdefault(int(ref[2]), []).append((source_id,field,section,Path(ref[1]).name))
                        else:
                            doc_id = add_document(conn,section,value)
                            conn.execute("INSERT INTO product_documents(source_id,document_id,field_name) VALUES(?,?,?)", (source_id,doc_id,field))
                    count += 1
                source_rows += count
                log(f"[import] {country} {table}: {count} rows", flush=True)
            if long_path:
                with long_path.open(encoding="utf-8-sig") as file:
                    for line_number, line in enumerate(file, 1):
                        if not line.strip():
                            continue
                        row = json.loads(line)
                        for source_id,field,section,filename in pending.pop(line_number, []):
                            if filename.casefold() != long_path.name.casefold():
                                raise ValueError(f"Wrong JSONL source for line {line_number}: {filename}")
                            if row.get("column") != field:
                                raise ValueError(f"JSONL column mismatch at line {line_number}")
                            doc_id = add_document(conn,section,row.get("value"))
                            conn.execute("INSERT OR REPLACE INTO product_documents VALUES(?,?,?,?,?)", (source_id,doc_id,field,str(long_path),line_number))
                        # Some exports omit content from the workbook but retain
                        # its ctid. Resolve those rows without relying on order.
                        if row.get("table") == "sr_product_description":
                            for source_id in ctid_links.get(str(row.get("ctid") or ""), []):
                                linked = conn.execute("SELECT 1 FROM product_documents WHERE source_id=? AND field_name='content'", (source_id,)).fetchone()
                                if not linked:
                                    doc_id = add_document(conn,"description",row.get("value"))
                                    conn.execute("INSERT INTO product_documents VALUES(?,?,?,?,?)", (source_id,doc_id,"content",str(long_path),line_number))
                        if line_number % 500 == 0:
                            log(f"[jsonl] {country} line={line_number}", flush=True)
            unresolved = sum(len(values) for values in pending.values())
            orphan_ids = [row[0] for row in conn.execute("""SELECT id FROM documents d
                WHERE NOT EXISTS (SELECT 1 FROM product_documents l WHERE l.document_id=d.id)""")]
            for document_id in orphan_ids:
                conn.execute("DELETE FROM document_search WHERE rowid=?", (document_id,))
                conn.execute("DELETE FROM spec_values WHERE document_id=?", (document_id,))
                conn.execute("DELETE FROM documents WHERE id=?", (document_id,))
            total = conn.execute("SELECT COUNT(*) FROM products WHERE country=?", (country,)).fetchone()[0]
            conn.execute("INSERT OR REPLACE INTO imports VALUES(?,?,?,?,?,?,?)", (country,str(workbook_path),str(long_path) if long_path else None,datetime.now(timezone.utc).isoformat(),total,source_rows,unresolved))
            conn.commit()
        except Exception:
            conn.rollback()
            raise
        finally:
            workbook.close()
        conn.execute("PRAGMA wal_checkpoint(TRUNCATE)")
    return {"country":country,"products":total,"source_rows":source_rows,"unresolved_refs":unresolved}
